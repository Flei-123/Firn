#!/usr/bin/env python3
# SPDX-License-Identifier: MPL-2.0
import struct
import sys
import zlib

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

LOREM = ("Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore "
         "et dolore magna aliqua. Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris nisi ut "
         "aliquip ex ea commodo consequat. Duis aute irure dolor in reprehenderit in voluptate velit esse "
         "cillum dolore eu fugiat nulla pariatur.")


def png_bytes(w, h, rgb):
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

    def chunk(tag, data):
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def add_page_field(par):
    run = par.add_run()
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def f01_plain(path):
    d = docx.Document()
    d.add_paragraph("Hello Firn. This is a plain paragraph.")
    d.add_paragraph(LOREM)
    d.add_paragraph("")
    d.add_paragraph("Third paragraph after an empty one.")
    d.save(path)


def f02_headings(path):
    d = docx.Document()
    d.add_heading("Title heading", 0)
    d.add_heading("Chapter one", 1)
    d.add_paragraph(LOREM)
    d.add_heading("Section 1.1", 2)
    d.add_paragraph(LOREM)
    d.add_heading("Section 1.2", 2)
    d.add_paragraph("Short.")
    d.save(path)


def f03_runs(path):
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Bold ").bold = True
    p.add_run("italic ").italic = True
    p.add_run("underline ").underline = True
    r = p.add_run("red ")
    r.font.color.rgb = RGBColor(255, 0, 0)
    r = p.add_run("big ")
    r.font.size = Pt(24)
    r = p.add_run("mono ")
    r.font.name = "Courier New"
    r = p.add_run("strike ")
    r.font.strike = True
    r = p.add_run("sup")
    r.font.superscript = True
    r = p.add_run("sub")
    r.font.subscript = True
    r = p.add_run(" caps text")
    r.font.all_caps = True
    r = p.add_run(" serif Times")
    r.font.name = "Times New Roman"
    d.save(path)


def f04_lists(path):
    d = docx.Document()
    d.add_paragraph("Bullets:")
    for t in ("alpha", "beta", "gamma"):
        d.add_paragraph(t, style="List Bullet")
    d.add_paragraph("Numbers:")
    for t in ("first", "second", "third"):
        d.add_paragraph(t, style="List Number")
    d.add_paragraph("nested", style="List Bullet 2")
    d.save(path)


def f05_table(path):
    d = docx.Document()
    d.add_paragraph("Before table")
    t = d.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    for r in range(4):
        for c in range(3):
            t.cell(r, c).text = "cell %d-%d text" % (r, c)
    d.add_paragraph("After table")
    d.save(path)


def f06_merged(path):
    d = docx.Document()
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    a = t.cell(0, 0).merge(t.cell(0, 2))
    a.text = "merged across three columns"
    b = t.cell(1, 0).merge(t.cell(2, 0))
    b.text = "vertical merge"
    t.cell(1, 1).text = "x1"
    t.cell(1, 2).text = "x2"
    t.cell(2, 1).text = "y1"
    t.cell(2, 2).text = "y2"
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), "FFFF00")
    t.cell(1, 1)._tc.get_or_add_tcPr().append(sh)
    d.save(path)


def f07_images(path, tmp="/tmp"):
    import os
    d = docx.Document()
    d.add_paragraph("Image below")
    for i, (w, h, col) in enumerate(((60, 40, (255, 0, 0)), (30, 30, (0, 160, 0)))):
        fp = os.path.join(tmp, "docx_fx_img%d.png" % i)
        with open(fp, "wb") as f:
            f.write(png_bytes(w, h, col))
        d.add_picture(fp, width=Cm(4 + 2 * i))
        os.remove(fp)
    d.add_paragraph("Text after the images.")
    d.save(path)


def f08_headerfooter(path):
    d = docx.Document()
    s = d.sections[0]
    s.header.paragraphs[0].text = "Header words here"
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ")
    add_page_field(fp)
    for i in range(80):
        d.add_paragraph("Body line %d: %s" % (i, LOREM[:90]))
    d.save(path)


def f09_pages(path):
    d = docx.Document()
    for i in range(6):
        d.add_heading("Chapter %d" % i, 1)
        for j in range(5):
            d.add_paragraph(LOREM)
        if i % 2 == 0:
            d.add_page_break()
    d.save(path)


def f10_align(path):
    d = docx.Document()
    for name, al in (("left", WD_ALIGN_PARAGRAPH.LEFT), ("center", WD_ALIGN_PARAGRAPH.CENTER),
                     ("right", WD_ALIGN_PARAGRAPH.RIGHT), ("justify", WD_ALIGN_PARAGRAPH.JUSTIFY)):
        p = d.add_paragraph("%s: %s" % (name, LOREM))
        p.alignment = al
    p = d.add_paragraph("Indented paragraph with first line indent. " + LOREM)
    p.paragraph_format.left_indent = Cm(2)
    p.paragraph_format.first_line_indent = Cm(1)
    p.paragraph_format.line_spacing = 1.5
    p = d.add_paragraph("tab\tone\ttwo\tthree")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(6), WD_TAB_ALIGNMENT.LEFT)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(10), WD_TAB_ALIGNMENT.RIGHT)
    p = d.add_paragraph("Line one")
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run("Line two after break")
    d.save(path)


def f11_sections(path):
    d = docx.Document()
    d.add_paragraph("Portrait section text. " + LOREM)
    s = d.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    d.add_paragraph("Landscape section text. " + LOREM)
    d.save(path)


def f12_nested(path):
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text = "outer a"
    t.cell(0, 1).text = "outer b"
    inner = t.cell(1, 0).add_table(rows=2, cols=2)
    inner.style = "Table Grid"
    for r in range(2):
        for c in range(2):
            inner.cell(r, c).text = "in%d%d" % (r, c)
    t.cell(1, 1).text = LOREM
    d.save(path)


def f13_unicode(path):
    d = docx.Document()
    d.add_paragraph("Umlaute: Äpfel Öl Übung ß € – “quotes” … ©")
    d.add_paragraph("Greek: αβγδ Ωμέγα  Cyrillic: Привет мир")
    d.add_paragraph("CJK: 你好世界 こんにちは 안녕")
    d.add_paragraph("Emoji: \U0001F600 \U0001F680 done")
    d.save(path)


def f14_longword(path):
    d = docx.Document()
    d.add_paragraph("x" * 400)
    d.add_paragraph("A" * 60 + " normal words follow " + "B" * 90)
    d.save(path)


def f15_empty(path):
    d = docx.Document()
    d.save(path)


def f16_longtable(path):
    d = docx.Document()
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Header A"
    t.rows[0].cells[1].text = "Header B"
    trpr = t.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))
    for i in range(70):
        r = t.add_row()
        r.cells[0].text = "row %d" % i
        r.cells[1].text = LOREM[:60]
    d.save(path)


FIXTURES = [f01_plain, f02_headings, f03_runs, f04_lists, f05_table, f06_merged, f07_images, f08_headerfooter,
            f09_pages, f10_align, f11_sections, f12_nested, f13_unicode, f14_longword, f15_empty, f16_longtable]


def main():
    out = sys.argv[1]
    for fn in FIXTURES:
        fn("%s/%s.docx" % (out, fn.__name__))


if __name__ == "__main__":
    main()
